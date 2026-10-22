#!/usr/bin/env python3
"""Single outreach file for the 100 live homepage concepts: output/OUTREACH_100.xlsx (one row per business with every
contact channel and the email drafts), plus OUTREACH_100.csv with the same columns.
Inputs: data/outreach_base.json (tools/outreach_base.py) and data/outreach_drafts_*.json (the drafts).
Usage: build_outreach.py"""
import csv, glob, json, re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

R = "/home/user/100homeimprovementsearch"
COLS = [("n", "#", 5), ("business_name", "Business", 26), ("owner", "Owner / contact", 20), ("owner_role", "Role", 16),
        ("old_site", "Current website", 30), ("new_concept", "New homepage concept", 40),
        ("google_phone", "Phone (Google profile)", 16), ("other_phones", "Other phones on their site", 22),
        ("whatsapp", "WhatsApp", 26), ("email", "Best email", 30), ("other_emails", "Other emails", 26),
        ("email_source", "Email found on", 30), ("contact_form_url", "Contact form", 30), ("facebook", "Facebook", 34),
        ("instagram", "Instagram", 34), ("other_socials", "Other profiles (Houzz, LinkedIn, YouTube...)", 40),
        ("address", "Address", 30), ("google_rating", "Google rating", 12), ("google_profile", "Google profile", 30),
        ("key_issue", "Main issue on current site", 50), ("subject", "Email subject", 34), ("email_body", "Email (day 1)", 70),
        ("follow_up_1", "Follow-up 1 (day 4)", 55), ("follow_up_2", "Follow-up 2 (day 11)", 50),
        ("short_dm", "Short message (Facebook / Instagram / SMS / WhatsApp)", 50), ("personalization_notes", "Why this angle", 40),
        ("check", "Check before sending", 50)]


# points the draft writers flagged for a human to confirm before sending (from their reports, 2026-10-22)
CHECK = {
    1: "The only email is asimkins@, an assistant's address found on the careers page; the email greets Matt. Consider calling the showroom or using the contact form first.",
    15: "The email goes to amy@ but greets Abe; Amy is likely office staff. Fine to send, or reach Abe by phone.",
    47: "The only email is ihor@ but the owner is Viktor; reviews mention Ihor or Igor, who may be family or a partner. Confirm who reads it.",
    49: "The only email is jennifer@ but the email greets Matt; consider the contact form or a call to reach Matt directly.",
    14: "Owner name: the record says Taha, but the email (talha@) and reviews say Talha. Confirm the spelling.",
    16: "Weaker fit: the firm has in-house marketing staff. Its site phone differs from Google; it may be a call-tracking number on purpose.",
    17: "The email mentions their October 17 showcase still promoted on the home page. Check it is still there before sending.",
    23: "Greeting uses Don (Donald in the record), as in his Google reviews.",
    28: "Follow-up 1 mentions a seminar block dated October 15 on their home page. Re-check before sending it.",
    36: "The email points out the giveaway banner says November 4 while the rules close entries November 1, 2026. Send before November 1 or remove that line.",
    38: "Greeting is to David, the current president; the record lists founder Yeheskel Sharbani as owner.",
    48: "The email mentions a 404 on their Full Home Remodels button (seen 2026-10-17). Re-check before sending.",
    57: "The new concept shows a 'Women-owned business' badge; the owner on record is Evan Hill. Confirm or remove the badge on the concept before sending.",
    58: "Their Google profile category is 'Home goods store' and the address starts with 'Parking lot'; follow-up 2 mentions this gently.",
    71: "Google and their site show different phone numbers; the concept uses Google's and the email asks which is right.",
    73: "The concept uses Google's (760) 642-6051, a California number, for a Durham NC firm; their site header shows (919) 289-5462. Confirm the main line.",
    76: "Owner on record is Michael Wood under new ownership since 2023; confirm he is the right contact.",
    77: "The observation (conflicting review counts on their site vs 24 on Google) is factual but the most sensitive; read it once before sending.",
    82: "The failing hero video was seen in a crawl and may play in a normal browser; the email words it as 'in a recent check'.",
    89: "Their Google business name is keyword-stuffed; the drafts use 'North2South'.",
    96: "Leads with a call link that dials 123456789 (verified). The contact form 404 may only affect GET requests, so the email suggests a test instead of claiming it is broken.",
    99: "Concept subdomain is pereiraconstructionsouthfl (their live domain); the record id is pcsouthflorida.com.",
}

def check_note(n, b, d):
    notes = [CHECK[n]] if n in CHECK else []
    best = b["verified_email"] or (b["emails"][0] if b["emails"] else "")
    if not best:
        notes.append("No email on file: use the contact form, then the short message on Facebook or Instagram.")
    if b.get("email_source") and "career" in b["email_source"]:
        notes.append("This email was found on their careers page and may reach staff, not the owner.")
    return " ".join(notes)


def main():
    base = {r["n"]: r for r in json.load(open(f"{R}/data/outreach_base.json"))}
    drafts = {}
    for f in sorted(glob.glob(f"{R}/data/outreach_drafts_*.json")):
        for d in json.load(open(f)):
            drafts[int(d["n"])] = d
    rows = []
    for n in sorted(base):
        b, d = base[n], drafts.get(n, {})
        s = b["socials"]
        best = b["verified_email"] or (b["emails"][0] if b["emails"] else "")
        cd = (b["conversion_notes"] or {}).get("current_site_defects") or []
        rows.append({
            "n": n, "business_name": b["business_name"], "owner": b["owner"] or "", "owner_role": b["owner_role"] or "",
            "old_site": b["old_site"], "new_concept": b["new_concept"], "google_phone": b["google_phone"] or "",
            "other_phones": ", ".join(p for p in b["all_phones"] if p != b["google_phone"]),
            "whatsapp": ", ".join(b["whatsapp_links_on_site"]) or "Not listed on their site (try the phone number on WhatsApp)",
            "email": best, "other_emails": ", ".join(e for e in b["emails"] if e != best),
            "email_source": b["email_source"] or "", "contact_form_url": b["contact_form_url"] or (b["old_site"] if not best else ""),
            "facebook": "\n".join(s.get("facebook", [])), "instagram": "\n".join(s.get("instagram", [])),
            "other_socials": "\n".join(u for k, v in s.items() if k not in ("facebook", "instagram") for u in v),
            "address": b["address"] or "", "google_rating": b["google_rating"], "google_profile": b["google_profile"] or "",
            "key_issue": "; ".join(b["measured_defects"]) or (cd[0] if cd else ""),
            **{k: d.get(k, "") for k in ("subject", "email_body", "follow_up_1", "follow_up_2", "short_dm", "personalization_notes")},
            "check": check_note(n, b, d)})

    wb = Workbook(); ws = wb.active; ws.title = "Outreach"
    ws.append([c[1] for c in COLS])
    for r in rows:
        ws.append([r[c[0]] for c in COLS])
    head = PatternFill("solid", fgColor="2B2420")
    for i, (_, _, w) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(1, i); c.font = Font(bold=True, color="FFFFFF"); c.fill = head; c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions

    g = wb.create_sheet("How to use")
    for line in [
        "How to use this file",
        "1. One row per business. 'New homepage concept' is the free preview you built; 'Current website' is their live site.",
        "2. Replace [Your Name], [Your Company], [Your Phone / WhatsApp] and [Your Business Address] in every email before sending.",
        "   US law (CAN-SPAM) requires a real postal address and a working opt-out in commercial email; the drafts already include the opt-out line.",
        "3. Send the day-1 email, then follow-up 1 after 3 days and follow-up 2 after 7 more days, only if there is no reply. Stop as soon as anyone says no.",
        "4. No email? Use the contact form link, then the short message on Facebook or Instagram.",
        "5. WhatsApp: none of these businesses links WhatsApp on its own site. The short message can be sent by SMS or WhatsApp to the Google phone,",
        "   but most US business lines are not on WhatsApp, so a phone call or SMS is the safer second channel.",
        "6. Send in small batches (10 to 20 a day) from a warmed-up mailbox to protect deliverability; personalised plain-text emails work best.",
        "7. Emails marked 'Email found on: careers page' reach staff, not the owner; consider the contact form or a call first for those.",
    ]:
        g.append([line])
    g.column_dimensions["A"].width = 140
    g["A1"].font = Font(bold=True, size=14)
    wb.save(f"{R}/output/OUTREACH_100.xlsx")

    with open(f"{R}/output/OUTREACH_100.csv", "w", newline="") as f:
        w = csv.writer(f); w.writerow([c[1] for c in COLS])
        for r in rows:
            w.writerow([r[c[0]] for c in COLS])
    print("rows", len(rows), "with drafts", sum(1 for r in rows if r["email_body"]))


if __name__ == "__main__":
    main()
